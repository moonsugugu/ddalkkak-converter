import hashlib
import io
import json
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import pymupdf
from PIL import Image
from docx import Document
from docx.shared import Inches
from pptx import Presentation

from app import load_settings
from converter import perform, targets, unavailable
from media import MediaCancelled, executable, probe


class UpgradeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.image = self.root/'원본 사진.jpg'
        Image.effect_noise((800, 600), 80).convert('RGB').save(self.image, quality=100, subsampling=0)
        self.original = hashlib.sha256(self.image.read_bytes()).hexdigest()

    def tearDown(self):
        self.tmp.cleanup()

    def test_f12_migration_and_saved_function_key(self):
        with patch('app.STATE_DIR', self.root):
            self.assertEqual(load_settings()['key'], 'F12')
            (self.root/'settings.json').write_text(json.dumps({'key':'Q', 'quality':'high'}))
            settings=load_settings()
            self.assertEqual((settings['version'],settings['key'],settings['quality']),(3,'F12','high'))
            (self.root/'settings.json').write_text(json.dumps({'version':2, 'key':'F9', 'quality':'small'}))
            self.assertEqual(load_settings()['key'], 'F9')

    def test_photo_formats_and_original_preserved(self):
        for target in ('png', 'webp', 'bmp', 'tiff', 'ico'):
            result = perform(self.image, target)
            with Image.open(result['files'][0]) as output:
                output.load()
                self.assertEqual(output.size, (256, 256) if target == 'ico' else (800, 600))
        self.assertEqual(hashlib.sha256(self.image.read_bytes()).hexdigest(), self.original)

    def test_photo_orientation_and_alpha(self):
        exif = Image.Exif()
        exif[274] = 6
        Image.new('RGB', (80, 50), 'red').save(self.image, exif=exif)
        with Image.open(perform(self.image, 'png')['files'][0]) as output:
            self.assertEqual(output.size, (50, 80))
        alpha = self.root/'투명.webp'
        Image.new('RGBA', (20, 30), (255, 0, 0, 0)).save(alpha, lossless=True)
        with Image.open(perform(alpha, 'png')['files'][0]) as output:
            self.assertEqual(output.convert('RGBA').getpixel((0, 0))[3], 0)

    def test_real_docx_and_pptx_photo_compression_and_extraction(self):
        docx = self.root/'사진 포함.docx'
        doc = Document()
        doc.add_paragraph('문서큐 사진 압축 검증')
        doc.add_picture(str(self.image), width=Inches(5))
        doc.save(docx)
        pptx = self.root/'사진 발표.pptx'
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = '사진 압축 검증'
        slide.shapes.add_picture(str(self.image), 0, 1000000, width=4000000)
        prs.save(pptx)
        for source in (docx, pptx):
            before = source.read_bytes()
            result = perform(source, 'compress')
            self.assertGreater(result['saved_percent'], 15)
            with zipfile.ZipFile(source) as original, zipfile.ZipFile(result['files'][0]) as compact:
                for name in original.namelist():
                    if name.endswith('.xml'):
                        self.assertEqual(original.read(name), compact.read(name))
            pictures = perform(result['files'][0], 'images')
            self.assertEqual(len(pictures['files']), 1)
            with Image.open(pictures['files'][0]) as image:
                self.assertEqual(image.size, (800, 600))
            self.assertEqual(source.read_bytes(), before)
        self.assertIn('문서큐', Document(perform(docx, 'compress')['files'][0]).paragraphs[0].text)
        self.assertEqual(len(Presentation(perform(pptx, 'compress')['files'][0]).slides), 1)

    def test_hwpx_media_only_rewrite_preserves_xml_and_mimetype(self):
        source = self.root/'한글 사진.hwpx'
        xml = b'<document><picture width="800" height="600"/><text>unchanged</text></document>'
        with zipfile.ZipFile(source, 'w') as out:
            out.writestr('mimetype', b'application/hwp+zip', compress_type=zipfile.ZIP_STORED)
            out.writestr('Contents/section0.xml', xml)
            out.writestr('BinData/image1.jpg', self.image.read_bytes())
        self.assertEqual(unavailable(source, 'compress', {}), '')
        result = perform(source, 'compress')
        self.assertGreater(result['saved_percent'], 15)
        with zipfile.ZipFile(result['files'][0]) as archive:
            self.assertEqual(archive.read('Contents/section0.xml'), xml)
            self.assertEqual(archive.getinfo('mimetype').compress_type, zipfile.ZIP_STORED)
            with Image.open(io.BytesIO(archive.read('BinData/image1.jpg'))) as picture:
                self.assertEqual(picture.size, (800, 600))
        self.assertEqual(len(perform(source, 'images')['files']), 1)

    def test_pdf_photo_compression_and_extraction_preserve_pages(self):
        source = self.root/'사진.pdf'
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((40, 40), 'Text remains editable')
        page.insert_image(pymupdf.Rect(40, 80, 540, 455), filename=str(self.image))
        doc.new_page().insert_text((40, 40), 'Second page')
        doc.save(source)
        doc.close()
        result = perform(source, 'compress')
        self.assertGreater(result['saved_percent'], 15)
        with pymupdf.open(result['files'][0]) as compact:
            self.assertEqual(len(compact), 2)
            self.assertIn('editable', compact[0].get_text())
        self.assertEqual(len(perform(source, 'images')['files']), 1)

    def test_no_savings_and_no_pictures_report_honestly(self):
        source = self.root/'작은.png'
        Image.new('RGB', (2, 2), 'white').save(source, optimize=True)
        result = perform(source, 'compress', quality='high')
        self.assertEqual(result['saved_percent'], 0)
        self.assertEqual(source.read_bytes(), Path(result['files'][0]).read_bytes())
        docx = self.root/'글만.docx'
        Document().save(docx)
        with self.assertRaisesRegex(ValueError, '사진이 없어요'):
            perform(docx, 'images')
        self.assertFalse(list(self.root.glob('.docq-*')))

    def test_new_target_availability(self):
        self.assertIn('png', targets('photo.jpg'))
        self.assertIn('png', targets('document.hwp'))
        self.assertIn('compress', targets('document.hwp'))
        self.assertIn('mp3', targets('movie.mp4'))
        self.assertIn('FFmpeg', unavailable('movie.mp4', 'mp3', {}))


@unittest.skipUnless(executable(), 'FFmpeg is not installed')
class MediaTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.tmp.name)
        cls.source = cls.root/'영상 예시.mp4'
        cmd = [executable(), '-hide_banner', '-loglevel', 'error', '-y', '-f', 'lavfi', '-i',
               'testsrc2=size=320x240:rate=12', '-f', 'lavfi', '-i', 'sine=frequency=440:sample_rate=44100',
               '-t', '1.5', '-c:v', 'libx264', '-preset', 'ultrafast', '-crf', '5',
               '-c:a', 'aac', '-threads', '2', str(cls.source)]
        subprocess.run(cmd, check=True, capture_output=True, creationflags=subprocess.CREATE_NO_WINDOW)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_mp4_video_gif_and_audio_outputs(self):
        before = self.source.read_bytes()
        for target in ('webm', 'mov', 'mkv', 'gif', 'mp3', 'wav'):
            with self.subTest(target=target):
                result = perform(self.source, target)
                output = Path(result['files'][0])
                self.assertGreater(output.stat().st_size, 0)
                if target == 'gif':
                    with Image.open(output) as image:
                        self.assertGreater(image.n_frames, 1)
                else:
                    info = probe(output)
                    self.assertGreater(float(info['format']['duration']), .5)
                    types = {s['codec_type'] for s in info['streams']}
                    self.assertIn('audio' if target in {'mp3','wav'} else 'video', types)
        self.assertEqual(self.source.read_bytes(), before)

    def test_video_compression_and_progress(self):
        progress = self.root/'progress.txt'
        result = perform(self.source, 'compress', progress_path=progress)
        self.assertGreater(result['saved_percent'], 20)
        self.assertIn('%', progress.read_text(encoding='utf-8'))
        self.assertLessEqual(result['result_bytes'], result['original_bytes'])

    def test_audio_transcoding(self):
        wav = Path(perform(self.source, 'wav')['files'][0])
        for target in ('m4a', 'flac', 'ogg'):
            result = perform(wav, target)
            self.assertEqual(probe(result['files'][0])['streams'][0]['codec_type'], 'audio')

    def test_cancellation_cleans_partial_work(self):
        cancel = self.root/'cancel'
        cancel.touch()
        with self.assertRaises(MediaCancelled):
            perform(self.source, 'webm', cancel_path=cancel)
        self.assertFalse(list(self.root.glob('.docq-*')))

    def test_silent_video_has_no_audio(self):
        source = self.root/'무음.mp4'
        subprocess.run([executable(), '-hide_banner', '-loglevel', 'error', '-y', '-i', str(self.source),
                        '-an', '-c:v', 'copy', str(source)], check=True, capture_output=True,
                       creationflags=subprocess.CREATE_NO_WINDOW)
        with self.assertRaisesRegex(ValueError, '소리가 없는'):
            perform(source, 'mp3')
        self.assertFalse(list(self.root.glob('.docq-*')))


if __name__ == '__main__':
    unittest.main()
