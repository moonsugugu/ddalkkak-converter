"""Installed-app check using generated documents; pump STA while workers run."""
import argparse
import json
import subprocess
import sys
import time
from pathlib import Path

import pythoncom
import win32com.client
import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from windows import process_running


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--exe', type=Path)
    args = parser.parse_args()
    folder = Path('validation/office').resolve()
    folder.mkdir(parents=True, exist_ok=True)
    cmd = [str(args.exe.resolve())] if args.exe else [sys.executable, str(Path('app.py').resolve())]
    report = {}

    def convert(path, target):
        print(f'Checking {path.suffix} → {target}', flush=True)
        request, response = folder / 'request.json', folder / 'response.json'
        request.write_text(json.dumps({'source':str(path), 'target':target,'quality':'balanced'}), encoding='utf-8')
        if response.exists():
            response.unlink()
        child = subprocess.Popen(cmd + ['--worker', str(request), str(response)],
                                 creationflags=subprocess.CREATE_NO_WINDOW)
        deadline = time.monotonic()+45
        while child.poll() is None and time.monotonic()<deadline:
            pythoncom.PumpWaitingMessages()
            time.sleep(.025)
        assert child.poll() == 0, 'worker failed or did not finish'
        result = json.loads(response.read_text(encoding='utf-8'))
        assert result['ok'], result
        files = [Path(p) for p in result['files']]
        assert files and all(p.stat().st_size > 0 for p in files)
        if target == 'pdf':
            with pymupdf.open(files[0]) as pdf:
                assert len(pdf)>0 and '문서큐' in pdf[0].get_text()
        report[path.suffix+' → '+target] = True
        return files[0]

    docx = folder / '문서큐 안내.docx'
    doc = Document()
    doc.add_heading('문서큐 변환 검증', 0)
    doc.add_paragraph('선택한 파일을 새로운 형식으로 바꿉니다.')
    doc.save(docx)
    sheet = folder / '문서큐 통계.xlsx'
    wb = Workbook()
    wb.active.append(['문서큐', '변환 수'])
    wb.active.append(['성공', 3])
    wb.save(sheet)
    slides = folder / '문서큐 발표.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = '문서큐 변환 검증'
    slide.placeholders[1].text = '한 번 눌러 변환합니다.'
    prs.save(slides)
    convert(docx, 'pdf')
    convert(sheet, 'pdf')
    convert(slides, 'pdf')
    hwp = convert(docx, 'hwp')
    convert(hwp, 'pdf')
    converted_docx = convert(hwp, 'docx')
    assert any('문서큐' in p.text for p in Document(converted_docx).paragraphs)
    hwpx = convert(hwp, 'hwpx')
    convert(hwpx, 'pdf')
    convert(hwp, 'xlsx')
    image_slides = convert(hwp, 'pptx')
    assert len(Presentation(image_slides).slides) > 0

    # HWP is a singleton: preserve an existing unsaved document while a worker runs.
    had_hwp = process_running('Hwp.exe')
    app = win32com.client.DispatchEx('HWPFrame.HwpObject')
    previous = app.XHwpDocuments.Active_XHwpDocument if had_hwp else None
    sentinel = app.XHwpDocuments.Add(False) if had_hwp else app.XHwpDocuments.Item(0)
    sentinel.SetActive_XHwpDocument()
    app.SetTextFile('DOCQ_UNSAVED_SENTINEL', 'TEXT', '')
    try:
        convert(hwp, 'pdf')
        sentinel.SetActive_XHwpDocument()
        assert 'DOCQ_UNSAVED_SENTINEL' in app.GetTextFile('TEXT', '')
        report['unsaved_hwp_preserved'] = True
    finally:
        sentinel.SetActive_XHwpDocument()
        sentinel.Close(False)
        if previous:
            previous.SetActive_XHwpDocument()
        if not had_hwp:
            try:
                app.Quit()
            except Exception:
                pass

    ppt = win32com.client.DispatchEx('PowerPoint.Application')
    sentinel = ppt.Presentations.Add(False)
    slide = sentinel.Slides.Add(1, 12)
    slide.Shapes.AddTextbox(1, 10, 10, 300, 40).TextFrame.TextRange.Text = 'DOCQ_UNSAVED_PPT'
    try:
        security = ppt.AutomationSecurity
        convert(slides, 'pdf')
        assert slide.Shapes(1).TextFrame.TextRange.Text == 'DOCQ_UNSAVED_PPT'
        assert ppt.AutomationSecurity == security
        report['unsaved_ppt_preserved'] = True
    finally:
        sentinel.Close()
        if ppt.Presentations.Count == 0:
            ppt.Quit()

    (folder/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
