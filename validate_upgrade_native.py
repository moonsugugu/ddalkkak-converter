"""Native Hancom checks through the real popup and generated photo documents."""
import hashlib
import json
import subprocess
import time
import tkinter as tk
from pathlib import Path

import win32api
import win32con
import win32event
import win32gui
from win32com.client import Dispatch
from PIL import Image, ImageGrab
from docx import Document
from docx.shared import Inches

from app import App
from validate_windows import focus_window, key


def main():
    folder = Path('validation/upgrade-native').resolve()
    folder.mkdir(parents=True, exist_ok=True)
    photo = folder/'큰 사진.jpg'
    Image.effect_noise((1600, 1200), 70).convert('RGB').save(photo, quality=100, subsampling=0)
    source = folder/'사진이 들어 있는 한글 예시.docx'
    doc = Document()
    doc.add_heading('문서큐 사진 압축 검증', 0)
    doc.add_paragraph('이 문장과 사진의 위치는 압축 후에도 유지합니다.')
    doc.add_picture(str(photo), width=Inches(5.5))
    doc.save(source)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    foreground = win32gui.GetForegroundWindow()
    root = tk.Tk()
    event = win32event.CreateEvent(None, True, False, None)
    app = App(root, event)
    app.settings['quality'] = 'balanced'
    report = {}
    def pump(seconds):
        end = time.monotonic()+seconds
        while time.monotonic()<end:
            root.update()
            time.sleep(.025)
    try:
        def convert(path, target):
            print(path.suffix+' -> '+target, flush=True)
            current = win32gui.GetForegroundWindow()
            if not current or win32gui.GetClassName(current) not in {'Progman', 'WorkerW'}:
                win32api.keybd_event(win32con.VK_LWIN, 0, 0, 0)
                key(ord('D'))
                win32api.keybd_event(win32con.VK_LWIN, 0, win32con.KEYEVENTF_KEYUP, 0)
                pump(.35)
            hwnd = win32gui.GetForegroundWindow()
            app.set_files([path], point=(750, 500), origin=hwnd)
            pump(.25)
            while target not in app.target_points:
                app.change_page(1)
                pump(.15)
            if target == 'compress':
                x, y = app.popup.winfo_rootx(), app.popup.winfo_rooty()
                ImageGrab.grab((x,y,x+460,y+460)).save(folder/'hwp-menu.png')
            x, y = app.target_points[target]
            win32api.SetCursorPos((app.popup.winfo_rootx()+round(x), app.popup.winfo_rooty()+round(y)))
            win32api.mouse_event(win32con.MOUSEEVENTF_LEFTDOWN,0,0)
            win32api.mouse_event(win32con.MOUSEEVENTF_LEFTUP,0,0)
            pump(.2)
            assert app.busy
            deadline=time.monotonic()+100
            while app.busy and time.monotonic()<deadline:
                pump(.1)
            assert not app.busy, 'Native conversion is waiting for Hancom'
            assert len(app.results)==1 and not app.errors, app.errors
            result=app.results[0]
            report[path.suffix+' -> '+target] = result
            assert not root.winfo_viewable() and not app.popup.winfo_viewable()
            return [Path(p) for p in result['files']]
        hwp = convert(source, 'hwp')[0]
        pages = convert(hwp, 'png')
        assert pages and all(Image.open(p).format == 'PNG' for p in pages)
        before = hashlib.sha256(hwp.read_bytes()).hexdigest()
        compact = convert(hwp, 'compress')[0]
        assert compact.stat().st_size < hwp.stat().st_size
        extracted = convert(compact, 'images')
        assert len(extracted) == 1
        with Image.open(extracted[0]) as image:
            assert image.size == (1600, 1200), image.size
        text = convert(compact, 'txt')[0].read_text(encoding='utf-8-sig')
        assert '문서큐 사진 압축 검증' in text
        compact_pages = convert(compact, 'png')
        assert len(compact_pages) == len(pages)
        assert Image.open(pages[0]).size == Image.open(compact_pages[0]).size
        assert hashlib.sha256(hwp.read_bytes()).hexdigest() == before
        assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
        report['originals_preserved'] = True
        report['embedded_photo_dimensions_preserved'] = True
        (folder/'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print('Native HWP PNG, photo compression, extraction and text preservation: PASS', flush=True)
    finally:
        if not app.busy:
            app.quit(force=True)
        win32api.CloseHandle(event)
        try:
            focus_window(foreground)
        except Exception:
            pass


if __name__ == '__main__':
    main()
