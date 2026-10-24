import csv
import hashlib
import io
import tempfile
import unittest
from pathlib import Path

import pymupdf
from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from converter import ConversionError, perform, read_text, targets, unavailable


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.pdf = self.root / '한글 이름.pdf'
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((50, 60), 'Document Q example')
        page.insert_text((50, 90), '=dangerous formula')
        pdf.new_page().insert_text((50, 60), 'Second page')
        pdf.save(self.pdf)
        pdf.close()

    def tearDown(self):
        self.temp.cleanup()

    def result(self, source, target):
        result = perform(source, target)
        return [Path(p) for p in result['files']]

    def test_pdf_images_pages_and_dpi(self):
        files = self.result(self.pdf, 'png')
        self.assertEqual(len(files), 2)
        with Image.open(files[0]) as image:
            self.assertEqual(image.size, (1488, 2105))

    def test_pdf_text_docx_xlsx_pptx(self):
        for target in ['txt', 'docx', 'xlsx', 'pptx']:
            with self.subTest(target=target):
                path = self.result(self.pdf, target)[0]
                if target == 'txt':
                    self.assertIn('Second page', path.read_text(encoding='utf-8-sig'))
                elif target == 'docx':
                    self.assertIn('Document Q', '\n'.join(p.text for p in Document(path).paragraphs))
                elif target == 'xlsx':
                    wb = load_workbook(path)
                    self.assertEqual(len(wb.worksheets), 2)
                    formula = next(c for row in wb.active for c in row
                                   if isinstance(c.value, str) and 'dangerous' in c.value)
                    self.assertEqual(formula.data_type, 's')
                    self.assertTrue(formula.value.startswith("'="))
                    wb.close()
                else:
                    prs = Presentation(path)
                    self.assertEqual(len(prs.slides), 2)
                    self.assertTrue(all(len(slide.shapes) == 1 for slide in prs.slides))

    def test_original_unchanged_and_collision_safe(self):
        before = hashlib.sha256(self.pdf.read_bytes()).hexdigest()
        a, b = self.result(self.pdf, 'txt'), self.result(self.pdf, 'txt')
        self.assertNotEqual(a[0].parent, b[0].parent)
        self.assertEqual(a[0].read_bytes(), b[0].read_bytes())
        self.assertEqual(hashlib.sha256(self.pdf.read_bytes()).hexdigest(), before)

    def test_scan_refuses_empty_text_and_cleans_up(self):
        scan = self.root / 'scan.pdf'
        pdf = pymupdf.open()
        pdf.new_page()
        pdf.save(scan)
        pdf.close()
        with self.assertRaisesRegex(ConversionError, 'OCR'):
            perform(scan, 'txt')
        self.assertFalse(list(self.root.glob('.docq-*')))

    def test_image_outputs_and_alpha(self):
        source = self.root / '투명.png'
        Image.new('RGBA', (100, 60), (255, 0, 0, 0)).save(source)
        for target in ['pdf', 'jpg', 'docx', 'pptx']:
            output = self.result(source, target)[0]
            self.assertGreater(output.stat().st_size, 0)
            if target == 'jpg':
                with Image.open(output) as image:
                    self.assertEqual(image.getpixel((0, 0)), (255, 255, 255))

    def test_multiframe_tiff_keeps_all_pages(self):
        source = self.root / 'multi.tiff'
        Image.new('RGB', (100, 200), 'red').save(source, save_all=True,
                                                append_images=[Image.new('RGB', (200, 100), 'blue')])
        self.assertEqual(len(self.result(source, 'png')), 2)
        with pymupdf.open(self.result(source, 'pdf')[0]) as pdf:
            self.assertEqual(len(pdf), 2)

    def test_sheets_csv_quotes_multiline_and_formulas(self):
        source = self.root / '통계.xlsx'
        wb = Workbook()
        wb.active.append(['학생', '내용', '수식'])
        wb.active.append(['문수', '첫째,둘째\n셋째', '=SUM(1,2)'])
        wb.create_sheet('둘째').append(['다른 시트'])
        wb.save(source)
        outputs = self.result(source, 'csv')
        self.assertEqual(len(outputs), 2)
        rows = list(csv.reader(io.StringIO(outputs[0].read_text(encoding='utf-8-sig'))))
        self.assertEqual(rows[1][1], '첫째,둘째\n셋째')
        self.assertEqual(rows[1][2], "'=SUM(1,2)")
        roundtrip = self.result(outputs[0], 'xlsx')[0]
        workbook = load_workbook(roundtrip)
        self.assertEqual(workbook.active['B2'].value, '첫째,둘째\n셋째')
        workbook.close()

    def test_encoding_and_txt_to_docx(self):
        source = self.root / '예시.txt'
        source.write_bytes('한글 문서입니다\n둘째 줄'.encode('cp949'))
        self.assertIn('한글', read_text(source))
        output = self.result(source, 'docx')[0]
        self.assertEqual(Document(output).paragraphs[0].text, '한글 문서입니다')
        source.write_bytes('한글'.encode('cp949'))
        self.assertEqual(read_text(source), '한글')

    def test_unsupported_and_capability_rules(self):
        self.assertNotIn('docx', targets('example.xlsx'))
        self.assertIn('pdf', targets('example.hwp'))
        self.assertIn('한글', unavailable('example.hwp', 'pdf', {}))
        with self.assertRaises(ConversionError):
            perform(self.pdf, 'hwp')
        self.assertFalse(list(self.root.glob('.docq-*')))


if __name__ == '__main__':
    unittest.main()
