"""Local text, ebook, subtitle, table and archive formats; never overwrite sources."""
import csv
import gzip
import html
import io
import re
import shutil
import stat
import tarfile
import textwrap
import uuid
import zipfile
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
import xml.etree.ElementTree as ET
from datetime import datetime,timezone

TEXT_INPUTS={'.md','.markdown','.html','.htm','.epub'}
ARCHIVES={'.zip','.tar','.gz','.tgz'}
EXTRA_INPUTS=TEXT_INPUTS|ARCHIVES|{'.srt','.vtt','.tsv','.ods'}
FORMAT_LABELS={'avif':'AVIF','md':'Markdown','html':'HTML','epub':'EPUB',
               'srt':'SRT','vtt':'VTT','tsv':'TSV','ods':'ODS'}
MAX_ARCHIVE_BYTES=512*1024*1024


def targets(ext):
    if ext in TEXT_INPUTS:values=['pdf','docx','txt','md','html','epub']
    elif ext in {'.srt','.vtt'}:values=['srt','vtt','txt']
    elif ext in {'.tsv','.ods'}:values=['xlsx','csv','tsv','txt','ods']
    elif ext in {'.txt','.docx','.pdf'}:values=['md','html','epub']
    elif ext in {'.csv','.xlsx'}:values=['tsv','ods']
    else:return []
    alias={'.markdown':'.md','.htm':'.html'}.get(ext,ext)
    return [value for value in values if '.'+value!=alias]


class TextHTML(HTMLParser):
    def __init__(self):super().__init__(convert_charrefs=True);self.parts=[];self.hidden=0
    def handle_starttag(self,tag,attrs):
        if tag in {'script','style','head'}:self.hidden+=1
        if not self.hidden and tag in {'p','div','br','li','h1','h2','h3','tr'}:self.parts.append('\n')
    def handle_endtag(self,tag):
        if tag in {'script','style','head'}:self.hidden=max(0,self.hidden-1)
        if not self.hidden and tag in {'p','div','li','h1','h2','h3','tr'}:self.parts.append('\n')
    def handle_data(self,data):
        if not self.hidden:self.parts.append(data)


def html_text(value):
    parser=TextHTML();parser.feed(value)
    return re.sub(r'\n[ \t]*\n[ \t]*\n+','\n\n',''.join(parser.parts)).strip()


def read_text(source):
    data=Path(source).read_bytes()
    if len(data)>16*1024*1024:raise ValueError('텍스트 파일은 16MB까지 지원해요.')
    for encoding in (('utf-16',) if data.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','cp949')):
        try:return data.decode(encoding)
        except UnicodeError:pass
    raise ValueError('텍스트 인코딩을 읽지 못했어요. UTF-8로 저장해 주세요.')


def archive_member(name):
    value=name.replace('\\','/')
    path=PurePosixPath(value)
    reserved={'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}
    if not value or path.is_absolute() or any(part in {'.','..'} or any(c in part for c in ':*?"<>|') or any(ord(c)<32 for c in part) or
            part.rstrip(' .')!=part or part.split('.')[0].upper() in reserved for part in path.parts):
        raise ValueError('안전하지 않은 압축 파일 경로가 있어 풀지 않았어요.')
    return path


def checked_zip(archive):
    infos=archive.infolist()
    if len(infos)>10000 or sum(i.file_size for i in infos)>MAX_ARCHIVE_BYTES:
        raise ValueError('압축 해제는 최대 10,000개·512MB까지 지원해요.')
    names=set()
    for info in infos:
        path=archive_member(info.filename.rstrip('/'))
        if stat.S_ISLNK(info.external_attr>>16):raise ValueError('링크가 들어 있는 압축 파일은 지원하지 않아요.')
        name=str(path).casefold()
        if name in names:raise ValueError('겹치는 압축 파일 이름이 있어요.')
        names.add(name)
    return infos


def document_text(source):
    source=Path(source);ext=source.suffix.lower()
    if ext=='.pdf':
        import pymupdf
        with pymupdf.open(source) as doc:
            if doc.needs_pass:raise ValueError('암호를 해제한 PDF 사본을 선택해 주세요.')
            return '\n\n'.join(page.get_text(sort=True) for page in doc)
    if ext=='.docx':
        from docx import Document
        doc=Document(source)
        return '\n'.join([p.text for p in doc.paragraphs]+['\t'.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    if ext=='.epub':
        with zipfile.ZipFile(source) as arc:
            checked_zip(arc)
            container=ET.fromstring(arc.read('META-INF/container.xml'))
            roots=container.findall('.//{*}rootfile')
            if not roots:raise ValueError('EPUB 목차를 읽지 못했어요.')
            opf=archive_member(roots[0].attrib['full-path'])
            package=ET.fromstring(arc.read(str(opf)))
            manifest={item.attrib['id']:item.attrib for item in package.findall('.//{*}manifest/{*}item')}
            chapters=[]
            for item in package.findall('.//{*}spine/{*}itemref'):
                entry=manifest.get(item.attrib.get('idref'),{})
                if entry.get('media-type')=='application/xhtml+xml':
                    path=archive_member(str(opf.parent/PurePosixPath(entry['href'])))
                    chapters.append(html_text(arc.read(str(path)).decode('utf-8-sig')))
            return '\n\n'.join(chapters)
    value=read_text(source)
    if ext in {'.html','.htm'}:return html_text(value)
    if ext in {'.md','.markdown'}:
        import markdown
        return html_text(markdown.markdown(value,extensions=['tables','fenced_code']))
    return value


def make_html(source,text):
    if source.suffix.lower() in {'.md','.markdown'}:
        import markdown
        class SafeHTML(HTMLParser):
            allowed={'p','h1','h2','h3','h4','h5','h6','ul','ol','li','strong','em','code','pre',
                     'blockquote','hr','br','table','thead','tbody','tr','th','td','del'}
            def __init__(self):super().__init__(convert_charrefs=True);self.parts=[]
            def handle_starttag(self,tag,attrs):
                if tag in self.allowed:self.parts.append('<'+tag+'>')
            def handle_endtag(self,tag):
                if tag in self.allowed:self.parts.append('</'+tag+'>')
            def handle_data(self,data):self.parts.append(html.escape(data))
        parser=SafeHTML()
        parser.feed(markdown.markdown(html.escape(read_text(source)),extensions=['tables','fenced_code']))
        rendered=''.join(parser.parts)
    else:rendered=''.join('<p>'+html.escape(line or ' ')+'</p>' for line in text.splitlines())
    return '<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>'+html.escape(source.stem)+'</title></head><body>'+rendered+'</body></html>'


def write_epub(path,title,text):
    title=html.escape(title);identifier='urn:uuid:'+str(uuid.uuid4())
    modified=datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    body=''.join('<p>'+html.escape(line or ' ')+'</p>' for line in text.splitlines())
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as arc:
        arc.writestr('mimetype','application/epub+zip',compress_type=zipfile.ZIP_STORED)
        arc.writestr('META-INF/container.xml','<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="EPUB/package.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        arc.writestr('EPUB/chapter.xhtml','<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" lang="ko"><head><title>'+title+'</title></head><body>'+body+'</body></html>')
        arc.writestr('EPUB/nav.xhtml','<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>목차</title></head><body><nav epub:type="toc"><ol><li><a href="chapter.xhtml">'+title+'</a></li></ol></nav></body></html>')
        arc.writestr('EPUB/package.opf','<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">'+identifier+'</dc:identifier><dc:title>'+title+'</dc:title><dc:language>ko</dc:language><meta property="dcterms:modified">'+modified+'</meta></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="chapter"/></spine></package>')


def text_job(source,work,target):
    text=document_text(source)
    if not text.strip():raise ValueError('내보낼 글자가 없어요. 스캔 문서는 먼저 OCR을 사용해 주세요.')
    if len(text)>1_000_000:raise ValueError('문서 글자가 너무 많아요. 나눠서 변환해 주세요.')
    destination=work/(source.stem+'.'+target)
    if target in {'txt','md'}:destination.write_text(text,encoding='utf-8-sig' if target=='txt' else 'utf-8')
    elif target=='html':destination.write_text(make_html(source,text),encoding='utf-8')
    elif target=='epub':write_epub(destination,source.stem,text)
    elif target=='docx':
        from docx import Document
        doc=Document()
        for line in text.splitlines():doc.add_paragraph(line)
        doc.save(destination)
    elif target=='pdf':
        import pymupdf
        lines=[wrapped for line in text.splitlines() for wrapped in (textwrap.wrap(line,43,replace_whitespace=False) or [''])]
        with pymupdf.open() as doc:
            for start in range(0,len(lines),45):
                page=doc.new_page();page.insert_text((40,50),'\n'.join(lines[start:start+45]),fontname='korea',fontsize=11,lineheight=1.5)
            doc.save(destination,deflate=True)
    else:raise ValueError('지원하지 않는 글 문서 형식이에요.')
    return {'text_only':True}


def subtitle_job(source,work,target):
    value=read_text(source).replace('\r\n','\n')
    blocks=re.split(r'\n\s*\n',value)
    cues=[]
    for block in blocks:
        lines=block.strip().splitlines()
        if not lines or lines[0].startswith(('WEBVTT','NOTE','STYLE','REGION')):continue
        position=next((i for i,line in enumerate(lines) if '-->' in line),None)
        if position is None:continue
        timing=lines[position]
        match=re.fullmatch(r'\s*((?:\d+:)?\d{2}:\d{2}[,.]\d{3})\s*-->\s*((?:\d+:)?\d{2}:\d{2}[,.]\d{3})(?:\s+.*)?',timing)
        if not match:raise ValueError('자막 시간 형식을 읽지 못했어요.')
        def normalized(stamp):
            stamp=stamp.replace(',','.')
            if stamp.count(':')==1:stamp='00:'+stamp
            parts=stamp.split(':')
            if int(parts[-2])>=60 or float(parts[-1])>=60:raise ValueError('자막 시간을 확인해 주세요.')
            return stamp
        start,end=map(normalized,match.groups())
        def seconds(stamp):
            h,m,s=stamp.split(':');return int(h)*3600+int(m)*60+float(s)
        if seconds(end)<=seconds(start):raise ValueError('자막 끝 시간이 시작보다 앞서 있어요.')
        cues.append((start,end,'\n'.join(lines[position+1:])))
    if not cues:raise ValueError('인식할 수 있는 자막이 없어요.')
    destination=work/(source.stem+'.'+target)
    if target=='txt':result='\n\n'.join(html_text(cue[2]) for cue in cues)
    elif target in {'srt','vtt'}:
        result='WEBVTT\n\n' if target=='vtt' else ''
        for i,(start,end,text) in enumerate(cues,1):
            if target=='srt':start,end=start.replace('.',','),end.replace('.',',')
            result+=f'{i}\n{start} --> {end}\n{text}\n\n'
    else:raise ValueError('지원하지 않는 자막 형식이에요.')
    destination.write_text(result,encoding='utf-8')
    return {'subtitle_cues':len(cues)}


def table_job(source,work,target):
    from converter import safe_cell
    from openpyxl import Workbook,load_workbook
    ext=source.suffix.lower()
    if ext=='.xlsx':
        book=load_workbook(source,read_only=True,data_only=False)
        try:rows=[list(row) for row in book.active.values]
        finally:book.close()
    elif ext=='.ods':
        with zipfile.ZipFile(source) as arc:
            checked_zip(arc);root=ET.fromstring(arc.read('content.xml'))
        rows=[]
        table=root.find('.//{*}table')
        if table is None:raise ValueError('ODS 시트를 읽지 못했어요.')
        for row in table.findall('{*}table-row'):
            values=[]
            for cell in row:
                repeat=int(next((v for k,v in cell.attrib.items() if k.endswith('number-columns-repeated')),'1'))
                if repeat>10000:raise ValueError('ODS 반복 셀 수가 너무 커요.')
                values.extend(['\n'.join(''.join(p.itertext()) for p in cell.findall('{*}p'))]*repeat)
            count=int(next((v for k,v in row.attrib.items() if k.endswith('number-rows-repeated')),'1'))
            if count>100000 or (len(rows)+count)*max(1,len(values))>1_000_000:raise ValueError('표가 너무 커요.')
            rows.extend([values]*count)
    else:rows=list(csv.reader(io.StringIO(read_text(source)),delimiter='\t' if ext=='.tsv' else ','))
    if sum(len(row) for row in rows)>1_000_000:raise ValueError('한 번에 최대 100만 셀을 변환해 주세요.')
    destination=work/(source.stem+'.'+target)
    if target=='xlsx':
        book=Workbook();sheet=book.active
        for row in rows:sheet.append([safe_cell(value) for value in row])
        book.save(destination)
    elif target=='ods':
        ns='urn:oasis:names:tc:opendocument:xmlns:'
        ET.register_namespace('office',ns+'office:1.0');ET.register_namespace('table',ns+'table:1.0');ET.register_namespace('text',ns+'text:1.0')
        root=ET.Element('{'+ns+'office:1.0}document-content',{'{'+ns+'office:1.0}version':'1.2'})
        table=ET.SubElement(ET.SubElement(ET.SubElement(root,'{'+ns+'office:1.0}body'),'{'+ns+'office:1.0}spreadsheet'),'{'+ns+'table:1.0}table',{'{'+ns+'table:1.0}name':'Sheet1'})
        for row in rows:
            row_element=ET.SubElement(table,'{'+ns+'table:1.0}table-row')
            for value in row:
                cell=ET.SubElement(row_element,'{'+ns+'table:1.0}table-cell',{'{'+ns+'office:1.0}value-type':'string'})
                ET.SubElement(cell,'{'+ns+'text:1.0}p').text='' if value is None else str(value)
        with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_DEFLATED) as arc:
            arc.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet',compress_type=zipfile.ZIP_STORED)
            arc.writestr('content.xml',ET.tostring(root,encoding='utf-8',xml_declaration=True))
            arc.writestr('META-INF/manifest.xml','<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    elif target in {'csv','tsv','txt'}:
        with destination.open('w',encoding='utf-8-sig',newline='') as stream:
            csv.writer(stream,delimiter=',' if target=='csv' else '\t').writerows(rows)
    else:raise ValueError('지원하지 않는 표 형식이에요.')
    return {'table_sheet':'첫 번째 시트','table_rows':len(rows)}


def archive_job(source,work,target,options,sources=None):
    paths=[Path(p).resolve(strict=True) for p in (sources or [source])]
    if target=='pack':
        fmt=options.get('format','zip')
        names=[p.name.casefold() for p in paths]
        if len(names)!=len(set(names)):raise ValueError('이름이 같은 파일은 따로 압축해 주세요.')
        for p in paths:
            if not p.is_file():raise ValueError('일반 파일만 압축할 수 있어요.')
            archive_member(p.name)
        destination=work/(source.stem+'.'+fmt)
        if fmt=='zip':
            with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_DEFLATED) as arc:
                for p in paths:arc.write(p,p.name)
        elif fmt=='tar':
            with tarfile.open(destination,'w') as arc:
                for p in paths:arc.add(p,arcname=p.name,recursive=False)
        elif fmt=='gz':
            if len(paths)!=1:raise ValueError('GZIP은 파일 한 개씩 압축해 주세요. 여러 개는 ZIP·TAR를 선택하세요.')
            destination=work/(source.name+'.gz')
            with source.open('rb') as incoming,gzip.open(destination,'wb') as output:shutil.copyfileobj(incoming,output)
        else:raise ValueError('지원하지 않는 압축 형식이에요.')
        return {'source_files':[str(p) for p in paths]}
    folder=work/(source.stem+'_풀기');folder.mkdir()
    if source.suffix.lower()=='.zip':
        with zipfile.ZipFile(source) as arc:
            infos=checked_zip(arc)
            for info in infos:
                destination=folder/Path(str(archive_member(info.filename.rstrip('/'))))
                if info.is_dir():destination.mkdir(parents=True,exist_ok=True)
                else:
                    destination.parent.mkdir(parents=True,exist_ok=True)
                    with arc.open(info) as incoming,destination.open('xb') as output:shutil.copyfileobj(incoming,output)
    elif source.suffix.lower() in {'.tar','.tgz'} or source.name.lower().endswith('.tar.gz'):
        with tarfile.open(source,'r:*') as arc:
            members=arc.getmembers()
            if len(members)>10000 or sum(m.size for m in members)>MAX_ARCHIVE_BYTES:raise ValueError('압축 해제 한도 10,000개·512MB를 넘었어요.')
            names=set()
            for member in members:
                path=archive_member(member.name.rstrip('/'));name=str(path).casefold()
                if name in names or not (member.isdir() or member.isfile()):raise ValueError('링크·특수 파일·중복 이름은 풀 수 없어요.')
                names.add(name)
            for member in members:
                destination=folder/Path(str(archive_member(member.name.rstrip('/'))))
                if member.isdir():destination.mkdir(parents=True,exist_ok=True)
                else:
                    destination.parent.mkdir(parents=True,exist_ok=True)
                    with arc.extractfile(member) as incoming,destination.open('xb') as output:shutil.copyfileobj(incoming,output)
    elif source.suffix.lower()=='.gz':
        name=archive_member(source.name[:-3])
        with gzip.open(source,'rb') as incoming,(folder/str(name)).open('xb') as output:
            total=0
            while chunk:=incoming.read(1024*1024):
                total+=len(chunk)
                if total>MAX_ARCHIVE_BYTES:raise ValueError('GZIP 해제 한도 512MB를 넘었어요.')
                output.write(chunk)
    else:raise ValueError('ZIP·TAR·GZIP 파일을 선택해 주세요.')
    return {'archive_extracted':True}
