"""Verify packaged worker modules and the external local media engine."""
import json
import subprocess
import time
from pathlib import Path

import pymupdf
from PIL import Image
from docx import Document
from docx.shared import Inches

from media import executable, probe


def main():
    folder = Path('validation/packaged-upgrade').resolve()
    folder.mkdir(parents=True, exist_ok=True)
    exe = Path('dist/DdalKkak.exe').resolve()
    photo = folder/'사진.jpg'
    Image.effect_noise((1200, 800), 75).convert('RGB').save(photo, quality=100)
    source = folder/'사진 문서.docx'
    doc = Document()
    doc.add_paragraph('문서큐 패키지 사진 압축 검사')
    doc.add_picture(str(photo), width=Inches(5))
    doc.save(source)
    pdf_source = folder/'사진.pdf'
    pdf = pymupdf.open()
    pdf.new_page().insert_image(pymupdf.Rect(50,50,550,450), filename=str(photo))
    pdf.save(pdf_source)
    pdf.close()
    video = folder/'영상 예시.mp4'
    subprocess.run([executable(), '-hide_banner','-loglevel','error','-y','-f','lavfi','-i',
                    'testsrc2=size=320x240:rate=12','-f','lavfi','-i','sine=frequency=440',
                    '-t','2','-c:v','libx264','-preset','ultrafast','-crf','5','-c:a','aac',
                    '-threads','2',str(video)], check=True, capture_output=True,
                   creationflags=subprocess.CREATE_NO_WINDOW)
    report = {}
    for path, target in [(source,'compress'), (pdf_source,'compress'), (video,'webm'), (video,'mp3'), (video,'gif')]:
        request, response = folder/'request.json', folder/'response.json'
        response.unlink(missing_ok=True)
        request.write_text(json.dumps({'source':str(path),'target':target,'quality':'balanced'}), encoding='utf-8')
        child = subprocess.Popen([str(exe),'--worker',str(request),str(response)],
                                 creationflags=subprocess.CREATE_NO_WINDOW)
        child.wait(timeout=60)
        result = json.loads(response.read_text(encoding='utf-8'))
        assert child.returncode == 0 and result['ok'], result
        if target == 'compress':
            assert result['saved_percent'] > 10, result
            if path.suffix == '.docx':
                assert '사진 압축 검사' in Document(result['files'][0]).paragraphs[0].text
            else:
                with pymupdf.open(result['files'][0]) as pdf:
                    assert len(pdf) == 1
        elif target == 'gif':
            with Image.open(result['files'][0]) as image:
                assert image.n_frames > 1
        else:
            assert float(probe(result['files'][0])['format']['duration']) > 1
        report[path.suffix+' -> '+target] = result
        print(path.suffix+' -> '+target+': PASS', flush=True)
    (folder/'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
